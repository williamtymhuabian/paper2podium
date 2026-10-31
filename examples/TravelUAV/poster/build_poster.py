#!/usr/bin/env python3
"""TravelUAV · A0 学术海报（research-ppt 示例）。

版式：A0 竖版 33.11 × 46.81 in，三栏，白底 + 学术蓝 + 红色强调。
数据与图均来自 Wang et al., arXiv:2410.07087。
"""
import math
import os

import qrcode
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(BASE, "..", "figures")
OUT = os.path.join(BASE, "TravelUAV_A0_海报.pptx")

W, H, M = 33.11, 46.81, 1.1
COL_W, GAP = 9.87, 0.65
XS = [M, M + COL_W + GAP, M + 2 * (COL_W + GAP)]

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xFB, 0xFA, 0xF7)
INK = RGBColor(0x1A, 0x1D, 0x24)
MUTED = RGBColor(0x5F, 0x66, 0x73)
PRIMARY = RGBColor(0x1F, 0x4E, 0x79)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)
LINE = RGBColor(0xD5, 0xDB, 0xE3)
SOFT = RGBColor(0xF2, 0xF5, 0xF8)

FONT = "Arial"
SERIF = "Georgia"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def rect(x, y, w, h, color, line=None, weight=1, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(weight)
    s.shadow.inherit = False
    return s


def rule(x, y, w, color=LINE, weight=1.2):
    s = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    s.line.color.rgb = color
    s.line.width = Pt(weight)


def text(x, y, w, h, paras, size=30, color=INK, font=FONT, bold=False,
         line=1.22, space_after=6, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    if isinstance(paras, str):
        paras = [[(paras, {})]]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line
        para.space_after = Pt(space_after)
        runs = p if isinstance(p, list) else [(p, {})]
        for t, st in runs:
            r = para.add_run()
            r.text = t
            r.font.size = Pt(st.get("size", size))
            r.font.color.rgb = st.get("color", color)
            r.font.bold = st.get("bold", bold)
            r.font.italic = st.get("italic", italic)
            r.font.name = st.get("font", font)
    return box


def section(x, y, num, title, w=COL_W):
    text(x, y, 0.8, 0.5, num, size=22, color=ACCENT, bold=True, font=FONT)
    text(x + 0.75, y - 0.06, w - 0.75, 0.75, title, size=40, color=PRIMARY, bold=True)
    rule(x, y + 0.78, w, PRIMARY, 1.6)
    return y + 1.0


def fig(path, x, y, w, caption):
    im = Image.open(path)
    h = w * im.height / im.width
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    cap_y = y + h + 0.14
    lines = max(1, math.ceil(len(caption) / 92))
    cap_h = lines * (19 * 1.15 / 72) + 0.1
    text(x, cap_y, w, cap_h, caption, size=19, color=MUTED, italic=True, line=1.15)
    return cap_y + cap_h + 0.18


def table(x, y, w, rows, col_w, header=True, size=21, row_h=0.62):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                 Inches(w), Inches(row_h * len(rows)))
    tbl = shp.table
    tbl.first_row = header
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if header and ri == 0:
                cell.fill.fore_color.rgb = PRIMARY
            else:
                cell.fill.fore_color.rgb = WHITE if ri % 2 else SOFT
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if ci else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(size)
            r.font.name = FONT
            r.font.bold = (header and ri == 0) or (str(val) == "Ours")
            r.font.color.rgb = WHITE if (header and ri == 0) else INK
    return y + row_h * len(rows)


def qr(path, data):
    q = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_H, box_size=16, border=4)
    q.add_data(data)
    q.make(fit=True)
    q.make_image(fill_color="#1F4E79", back_color="#FFFFFF").save(path)


# ---------------------------------------------------------------- 页眉
rect(0, 0, W, 0.32, PRIMARY)
text(M, 0.85, 23.0, 0.5, "UAV VISION-LANGUAGE NAVIGATION   ·   PREPRINT 2024",
     size=21, color=ACCENT, bold=True, font=FONT)
text(M, 1.42, 24.6, 1.6, "Towards Realistic UAV Vision-Language Navigation",
     size=72, color=INK, bold=True, line=1.02)
text(M, 3.55, 24.6, 0.9, "Platform, Benchmark, and Methodology",
     size=52, color=PRIMARY, bold=True, line=1.0)
text(M, 4.72, 24.8, 0.55,
     "Xiangyu Wang*, Donglin Yang*, Ziqin Wang*, Hohin Kwan, Jinyu Chen, "
     "Wenjun Wu, Hongsheng Li, Yue Liao\u2020, Si Liu\u2020",
     size=24, color=INK, line=1.15)
text(M, 5.32, 24.8, 0.9,
     "Institute of Artificial Intelligence, Beihang University  ·  MMLab, CUHK  ·  "
     "Centre for Perceptual and Interactive Intelligence",
     size=21, color=MUTED, line=1.2)
text(M, 5.95, 24.8, 0.4, "* Equal contribution      \u2020 Corresponding author",
     size=17, color=MUTED, italic=True)

qr(os.path.join(BASE, "qr_arxiv.png"), "https://arxiv.org/abs/2410.07087")
slide.shapes.add_picture(os.path.join(BASE, "qr_arxiv.png"),
                         Inches(28.35), Inches(1.15), Inches(3.1), Inches(3.1))
text(28.35, 4.4, 3.1, 0.9,
     [[("arXiv:2410.07087", {"bold": True, "color": INK})],
      [("Project page", {"color": MUTED})]],
     size=19, align=PP_ALIGN.CENTER, line=1.2)
text(28.35, 5.45, 3.1, 0.4, "prince687028.github.io/OpenUAV",
     size=15, color=MUTED, align=PP_ALIGN.CENTER)
rule(M, 6.55, W - 2 * M, PRIMARY, 2.2)

# ---------------------------------------------------------------- 第一栏
x = XS[0]
y = section(x, 7.0, "01", "MOTIVATION")
text(x, y, COL_W, 5.0, [
    [("UAV object search is still flown as waypoints. ", {"bold": True}),
     ("A human translates \u201cwhat to find\u201d into \u201chow to fly\u201d before every mission.", {})],
    [("Ground VLN assumes discrete moves on a graph. ", {"bold": True}),
     ("A drone moves continuously in 6 DoF, with inertia, attitude and obstacle constraints.", {})],
    [("Simplified simulators hide the gap. ", {"bold": True}),
     ("Policies that pass in simulation drift on real aircraft; new maps and targets demand fresh data.", {})],
], size=28, line=1.26, space_after=12)

rect(x, 12.55, COL_W, 1.35, SOFT)
rect(x, 12.55, 0.12, 1.35, ACCENT)
text(x + 0.35, 12.75, COL_W - 0.6, 1.0,
     "We need a platform, a benchmark, and a method built for continuous flight.",
     size=27, color=PRIMARY, bold=True, line=1.18)

y = section(x, 14.35, "02", "WHY BENCHMARKS FALL SHORT")
text(x, y, COL_W, 3.0, [
    [("Action space. ", {"bold": True}), ("2\u20134 DoF discrete actions cannot express continuous flight.", {})],
    [("Dynamics. ", {"bold": True}), ("Simplified kinematics hide sim-to-real drift.", {})],
    [("Scale. ", {"bold": True}), ("Few datasets cover many objects, scenes and instruction styles.", {})],
], size=27, line=1.24, space_after=9)

t1 = [
    ["Dataset", "DoF", "Kin.", "#Traj", "Vocab", "Actions", "Environment"],
    ["R2R", "\u2013", "\u2717", "7,189", "3.1K", "5", "Matterport3D"],
    ["TouchDown", "2", "\u2717", "9,326", "5K", "35", "Street View"],
    ["AerialVLN", "4", "\u2717", "8,446", "4.5K", "204", "AirSim"],
    ["CityNav", "4", "\u2717", "32,637", "6.6K", "\u2013", "SensatUrban"],
    ["Ours", "6", "\u2713", "12,149", "10.8K", "264", "OpenUAV"],
]
y = table(x, 18.75, COL_W, t1, [2.05, 0.75, 0.8, 1.45, 1.25, 1.35, 2.22], size=20, row_h=0.72) + 0.15
text(x, y, COL_W, 0.4, "Table 1. Existing VLN datasets vs. ours (6 DoF, real kinematics).",
     size=18, color=MUTED, italic=True)

y = section(x, 24.55, "03", "OPENUAV PLATFORM")
y = fig(os.path.join(FIG, "fig1_platform.png"), x, y, COL_W,
        "Figure 1. Closed-loop simulation, human control interface, configurable sensors and trajectory replay.")
text(x, y + 0.02, COL_W, 2.6, [
    [("Unreal Engine 4 + AirSim", {"bold": True}), (", 22 high-fidelity scenes with diverse objects and lighting.", {})],
    [("Closed-loop, parallelized simulation", {"bold": True}), (" with reproducible conditions.", {})],
    [("Human control interface", {"bold": True}), (" records continuous flights; sensors are collected asynchronously.", {})],
], size=26, line=1.24, space_after=8)

rule(x, 35.15, COL_W, ACCENT, 3)
text(x, 35.45, COL_W, 2.2,
     "One platform, one benchmark, one policy \u2014 all built for continuous 6-DoF flight.",
     size=34, color=PRIMARY, bold=True, line=1.18)

y = section(x, 38.35, "04", "CONTRIBUTIONS")
text(x, y, COL_W, 5.6, [
    [("1  OpenUAV. ", {"bold": True, "color": PRIMARY}),
     ("A realistic, closed-loop UAV simulation platform with a human control interface.", {})],
    [("2  UAV-Need-Help. ", {"bold": True, "color": PRIMARY}),
     ("A 12k-trajectory, 6-DoF benchmark with three levels of assistant guidance.", {})],
    [("3  Hierarchical navigation LLM. ", {"bold": True, "color": PRIMARY}),
     ("Far target pose, then fine trajectory, trained with backtracking data aggregation.", {})],
], size=26, line=1.24, space_after=12)

# ---------------------------------------------------------------- 第二栏
x = XS[1]
y = section(x, 7.0, "05", "DATA COLLECTION")
y = fig(os.path.join(FIG, "fig2_dataset.png"), x, y, COL_W,
        "Figure 2. (a) Instruction vocabulary, (b\u2013d) trajectory and target statistics, (e) split design.")
text(x, y + 0.02, COL_W, 3.0, [
    [("Human pilots", {"bold": True}), (" fly continuous trajectories; UAV states are recorded and sensor data is collected asynchronously.", {})],
    [("GPT-4 descriptions", {"bold": True}), (" are generated for each target and manually checked.", {})],
    [("Each trajectory ends within 5 m of the target", {"bold": True}), ("; collisions are discarded.", {})],
], size=26, line=1.24, space_after=8)

y = section(x, 19.0, "06", "DATASET AT A GLANCE")
stats = [("12,149", "trajectories"), ("6 DoF", "continuous flight"), ("89", "object categories"), ("22", "scenes")]
for i, (num, lab) in enumerate(stats):
    sx = x + (i % 2) * (COL_W / 2)
    sy = y + (i // 2) * 1.65
    text(sx, sy, COL_W / 2 - 0.2, 0.9, num, size=52, color=PRIMARY, bold=True, line=1.0)
    text(sx, sy + 0.78, COL_W / 2 - 0.2, 0.4, lab, size=22, color=MUTED)
y += 3.4
t2 = [
    ["Split", "Traj.", "Objects", "Scenes"],
    ["Train", "9,152", "76", "20"],
    ["Test Seen", "1,410", "seen", "seen"],
    ["Test Unseen Map", "958", "seen", "2 unseen"],
    ["Test Unseen Object", "629", "13 unseen", "seen"],
]
y = table(x, y, COL_W, t2, [3.4, 2.0, 2.2, 2.27], size=21, row_h=0.68) + 0.12
text(x, y, COL_W, 0.4,
     "Easy < 250 m  ·  Hard > 250 m  ·  target distance 50\u2013400 m",
     size=19, color=MUTED, italic=True)

y = section(x, 29.6, "07", "UAV-NEED-HELP BENCHMARK")
text(x, y, COL_W, 1.0,
     "The UAV searches from the target description, the scene, and guidance from an assistant.",
     size=25, color=INK, line=1.22)
levels = [
    ("L1", "Full guidance", "Continuous guidance to the target throughout the episode."),
    ("L2", "Coarse guidance", "Guidance only at critical states, when the policy is uncertain."),
    ("L3", "Avoidance", "Warns about collisions and hazards only."),
]
ly = y + 1.15
for num, title, desc in levels:
    rect(x, ly, 0.85, 0.85, PRIMARY)
    text(x, ly + 0.18, 0.85, 0.5, num, size=26, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(x + 1.1, ly - 0.05, COL_W - 1.1, 0.45, title, size=26, color=PRIMARY, bold=True)
    text(x + 1.1, ly + 0.42, COL_W - 1.1, 0.75, desc, size=22, color=INK, line=1.18)
    ly += 1.25

y = section(x, 37.0, "08", "WHAT THE INSTRUCTIONS LOOK LIKE")
rect(x, y, COL_W, 3.2, SOFT)
text(x + 0.35, y + 0.3, COL_W - 0.7, 2.6,
     "\u201cThe target is located to the front right of the drone. The target is a person "
     "wearing a brown jacket and blue pants. Nearby, there is a swimming pool and four chairs. "
     "Please control the drone and find the target.\u201d",
     size=24, color=INK, italic=True, line=1.3)
text(x + 0.35, y + 2.55, COL_W - 0.7, 0.4,
     "Target descriptions combine direction, appearance, and surrounding context.",
     size=19, color=MUTED)

# ---------------------------------------------------------------- 第三栏
x = XS[2]
y = section(x, 7.0, "09", "METHOD: HIERARCHICAL")
y = fig(os.path.join(FIG, "fig3_method.png"), x, y, COL_W,
        "Figure 3. Instruction and five views are encoded by an MLLM; a far target pose is decoded first, then a fine 6-DoF trajectory.")
text(x, y + 0.02, COL_W, 4.0, [
    [("Understand. ", {"bold": True}), ("An MLLM fuses the instruction with five RGB views and produces a task representation.", {})],
    [("Decide far, then near. ", {"bold": True}), ("The model predicts a far target pose, then a fine trajectory in continuous action space.", {})],
    [("Learn from backtracking. ", {"bold": True}), ("Data aggregation samples recovery trajectories, improving obstacle avoidance and robustness.", {})],
], size=25, line=1.24, space_after=9)

rect(x, 16.35, COL_W, 2.35, SOFT)
rect(x, 16.35, 0.12, 2.35, ACCENT)
text(x + 0.4, 16.6, COL_W - 0.8, 1.9, [
    [("Why far-to-near? ", {"bold": True, "color": PRIMARY}),
     ("Direct regression must cover a huge continuous search space. A far target pose narrows "
      "down where to go; the fine trajectory covers a short horizon, and the policy re-plans "
      "from new views.", {})],
], size=23, color=INK, line=1.24)

y = section(x, 20.2, "10", "RESULTS: TEST SEEN (L1)")
text(x, y - 0.05, COL_W, 1.2, "2.1\u00d7", size=150, color=ACCENT, bold=True, line=0.9)
text(x + 0.05, y + 1.98, COL_W, 0.9,
     "success rate over the CMA baseline", size=28, color=INK, bold=True)
text(x + 0.05, y + 2.5, COL_W, 0.5,
     "17.45% vs 8.37%  ·  same protocol", size=24, color=MUTED)
t3 = [
    ["Method", "NE \u2193", "SR \u2191", "OSR \u2191", "SPL \u2191"],
    ["CMA", "135.73", "8.37", "18.72", "7.90"],
    ["Ours", "106.28", "16.10", "44.26", "14.30"],
    ["Ours-DA", "98.66", "17.45", "48.87", "15.76"],
    ["Human", "14.15", "94.51", "94.51", "77.84"],
]
y = table(x, y + 3.2, COL_W, t3, [2.6, 1.75, 1.75, 1.75, 1.85], size=21, row_h=0.7) + 0.12
text(x, y, COL_W, 0.6,
     "Table 2. Test Seen, L1 assistant. Simulation results; real flights remain subject to airspace and site conditions.",
     size=18, color=MUTED, italic=True, line=1.18)

y = section(x, 31.4, "11", "GENERALIZATION")
text(x, y, COL_W / 2 - 0.25, 0.7, "22.42%", size=48, color=PRIMARY, bold=True)
text(x, y + 0.8, COL_W / 2 - 0.25, 0.9, "SR on unseen objects\n(CMA: 9.06%)", size=21, color=INK, line=1.22)
text(x + COL_W / 2 + 0.1, y, COL_W / 2 - 0.25, 0.7, "4.18%", size=48, color=PRIMARY, bold=True)
text(x + COL_W / 2 + 0.1, y + 0.8, COL_W / 2 - 0.25, 0.9, "SR on unseen maps\n(CMA: 2.30%)", size=21, color=INK, line=1.22)
text(x, y + 1.75, COL_W, 1.4,
     "The same policy transfers to new maps and new object categories without retraining.",
     size=25, color=INK, line=1.24)

y = fig(os.path.join(FIG, "fig4_qualitative.png"), x, 37.1, COL_W,
        "Figure 4. Closed-loop rollouts in held-out scenes; when the target is ambiguous, the policy re-observes.")

# ---------------------------------------------------------------- 页脚
rect(0, 45.55, W, 0.16, PRIMARY)
text(M, 45.85, W - 2 * M, 0.5,
     "Wang et al.  ·  Towards Realistic UAV Vision-Language Navigation  ·  arXiv:2410.07087  ·  "
     "Code, data and benchmark: prince687028.github.io/OpenUAV  ·  Demo poster, non-commercial use.",
     size=17, color=MUTED)

prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
